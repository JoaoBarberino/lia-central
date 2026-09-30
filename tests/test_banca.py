"""Correções que a banca simulada achou (30/09): datas escritas de outros jeitos, planilha com células
problemáticas, código em conflito, reabrir depois de um bloqueio e textos."""
import json
import os
import re
from datetime import datetime

import openpyxl

from central import activities as acts
from central.ai import FakeLLM, date_in_text, mentioned, validate_item
from central.db import connect

from .test_web import client, login  # noqa: F401  (fixture)


# ---------------------------------------------------------------------------
# Datas e nomes na checagem da IA
# ---------------------------------------------------------------------------
def test_datas_escritas_de_varios_jeitos():
    assert date_in_text("2026-10-07", "o prazo passou para 7 de outubro")          # o exemplo do próprio case
    assert date_in_text("2026-10-07", "até 07/10/2026") and date_in_text("2026-10-07", "dia 7/10.")
    assert date_in_text("2026-10-14", "quarta-feira, dia 14 de outubro")
    assert date_in_text("2026-10-01", "1º de outubro")
    assert not date_in_text("2026-10-07", "até 17/10")                             # 7/10 não está dentro de 17/10
    assert not date_in_text("2026-10-07", "27 de outubro")
    assert not date_in_text("2026-10-07", "07/10/2025")                            # outro ano
    assert not date_in_text("2026-10-16", "até sexta que vem")                    # data calculada não é fato


def test_nome_como_palavra_inteira():
    assert not mentioned("Ana", "fica para a próxima semana")
    assert mentioned("Ana", "a Ana, depois") and mentioned("Davi", "Davi assumiu")


def _item(**kw):
    return {"kind": "update", "reason": "teste", "uncertainties": []} | kw


def test_prazo_por_extenso_vira_sugestao(conn, sync):
    sync()
    texto = "O carrossel sobre ferramentas (ACT-101) continua com Ana; o prazo passou para 7 de outubro."
    clean, problema = validate_item(conn, _item(target_activity_id="ACT-101", due_date="2026-10-07",
                                                owners=["U-A"], evidence=texto), texto)
    assert problema is None and clean["proposed"] == {"due_date": "2026-10-07"}


def test_prazo_calculado_pede_conferencia_e_nao_diz_que_ja_esta_no_quadro(conn, sync):
    sync()
    texto = "Ficou combinado que o Davi entrega o checklist (ACT-102) até sexta que vem."
    clean, problema = validate_item(conn, _item(target_activity_id="ACT-102", due_date="2026-10-16",
                                                evidence=texto), texto)
    assert problema == "precisa_conferir" and "não aparece escrito como data" in clean["dropped"][0]


def test_prazo_calculado_aparece_como_precisa_de_conferencia(conn, sync, folder):
    sync()
    (folder / "Ata_2026-10-07.md").write_text(
        "# Ata de reunião de 7 de outubro de 2026\n\nstatus: ativo\ndata_da_reuniao: 2026-10-07\n\n"
        "Participaram Davi e Bruno.\n\nFicou combinado que o Davi entrega o checklist (ACT-102) até sexta que vem.\n",
        encoding="utf-8")
    llm = FakeLLM(lambda s, u: {"items": [_item(target_activity_id="ACT-102", due_date="2026-10-16",
                  evidence="Ficou combinado que o Davi entrega o checklist (ACT-102) até sexta que vem.")]})
    sync(llm=llm)
    nota = conn.execute("SELECT kind, reason FROM extraction_notes WHERE kind='precisa_conferir'").fetchone()
    assert nota and nota["reason"].startswith("ACT-102:")
    assert not conn.execute("SELECT 1 FROM extraction_notes WHERE kind='sem_mudanca' "
                            "AND text LIKE '%sexta que vem%'").fetchone()
    assert not conn.execute("SELECT 1 FROM suggestions WHERE target_activity_id='ACT-102'").fetchone()


def test_titulo_parecido_pede_conferencia(conn, sync):
    sync()
    texto = "Ficou decidido que Ana vai preparar carrossel de ferramentas até 2026-10-20."
    clean, problema = validate_item(conn, {"kind": "create", "title": "Preparar carrossel de ferramentas",
                                           "owners": ["U-A"], "due_date": "2026-10-20", "evidence": texto,
                                           "reason": "nova", "uncertainties": []}, texto)
    assert problema is None and any("Parecida com a atividade ACT-101" in u for u in clean["uncertainties"])


# ---------------------------------------------------------------------------
# Planilha oficial
# ---------------------------------------------------------------------------
def _editar_planilha(folder, mudar):
    wb = openpyxl.load_workbook(folder / "Ata_registro.xlsx")
    ws = wb.active
    cabecalho = [c.value for c in ws[1]]
    mudar(ws, cabecalho)
    wb.save(folder / "Ata_registro.xlsx")


def _col(cabecalho, nome):
    return cabecalho.index(nome) + 1


def test_prazo_digitado_como_texto_e_convertido(conn, sync, folder):
    sync()
    _editar_planilha(folder, lambda ws, h: setattr(ws.cell(row=3, column=_col(h, "Prazo")), "value", "09/10/2026"))
    sync()
    s = conn.execute("SELECT proposed_fields, evidence FROM suggestions WHERE target_activity_id='ACT-102'").fetchone()
    assert json.loads(s["proposed_fields"]) == {"due_date": "2026-10-09"}
    assert "'09/10/2026'" in s["evidence"]                                        # evidência = texto da célula


def test_prazo_ilegivel_nao_apaga_o_prazo_e_vira_pendencia(conn, sync, folder):
    sync()
    _editar_planilha(folder, lambda ws, h: setattr(ws.cell(row=3, column=_col(h, "Prazo")), "value", "sexta"))
    sync()
    assert not conn.execute("SELECT 1 FROM suggestions WHERE target_activity_id='ACT-102'").fetchone()
    assert conn.execute("SELECT 1 FROM issues WHERE title LIKE 'Prazo ilegível%' AND status='aberta'").fetchone()
    assert acts.snapshot(conn, "ACT-102")["due_date"] == "2026-10-06"


def test_carga_inicial_com_prazo_em_texto_importa_todas(conn, folder):
    from central.sources import LocalSource
    from central.sync import run_sync

    def mudar(ws, h):
        ws.cell(row=3, column=_col(h, "Prazo")).value = "06/10/2026"
        ws.cell(row=4, column=_col(h, "Prazo")).value = "quando der"
    _editar_planilha(folder, mudar)
    run_sync(conn, LocalSource(folder), llm=None)
    assert {a["activity_id"] for a in acts.list_activities(conn, include_done=True)} == \
        {"ACT-101", "ACT-102", "ACT-103", "ACT-104"}                               # nenhuma linha se perdeu
    assert acts.snapshot(conn, "ACT-102")["due_date"] == "2026-10-06"
    assert acts.snapshot(conn, "ACT-103")["due_date"] is None                     # "quando der" → a definir
    assert conn.execute("SELECT 1 FROM issues WHERE title LIKE 'Prazo ilegível%'").fetchone()


def test_responsavel_desconhecido_na_planilha_vira_pendencia(conn, sync, folder):
    sync()
    _editar_planilha(folder, lambda ws, h: setattr(ws.cell(row=5, column=_col(h, "Responsáveis")), "value", "Ana; Eduardo"))
    sync()
    assert conn.execute("SELECT 1 FROM issues WHERE kind='responsavel_desconhecido' AND detail LIKE '%Eduardo%'").fetchone()
    s = conn.execute("SELECT evidence, uncertainties FROM suggestions WHERE target_activity_id='ACT-104'").fetchone()
    assert "'Ana; Eduardo'" in s["evidence"] and "Eduardo" in s["uncertainties"]


def test_codigo_em_conflito_nao_sobrescreve_atividade_da_central(conn, sync, folder):
    sync()
    novo = acts.create_activity(conn, {"title": "Organizar planilha de presença", "owners": ["U-D"]},
                                actor_id="U-D", creation_kind="manual", reason="Criada na interface")
    assert novo == "ACT-105"
    _editar_planilha(folder, lambda ws, h: ws.append(["ACT-105", "Preparar oficina de prompts", "Carla"]))
    sync()
    assert not conn.execute("SELECT 1 FROM suggestions WHERE target_activity_id='ACT-105'").fetchone()
    assert conn.execute("SELECT 1 FROM issues WHERE title='Código ACT-105 em conflito'").fetchone()
    assert acts.snapshot(conn, "ACT-105")["title"] == "Organizar planilha de presença"


# ---------------------------------------------------------------------------
# Interface
# ---------------------------------------------------------------------------
def _db():
    return connect(os.environ["DATABASE_PATH"])


def test_reabrir_atividade_que_estava_bloqueada(client):  # noqa: F811
    login(client, "U-D")
    client.post("/atividades/ACT-102/estado", data={"status": "Bloqueada", "reason": "Esperando a secretaria"})
    client.post("/atividades/ACT-102/estado", data={"status": "Concluída"})
    pagina = client.get("/atividades/ACT-102").text
    valor = re.search(r'name="status" value="([^"]+)"><button[^>]*>Reabrir', pagina).group(1)
    client.post("/atividades/ACT-102/estado", data={"status": valor})
    a = acts.snapshot(_db(), "ACT-102")
    assert a["status"] == "Bloqueada"
    assert acts.block_reason(_db(), a | {"activity_id": "ACT-102"}) == "Esperando a secretaria"


def test_ia_desligada_avisa_na_tela_de_sugestoes(client):  # noqa: F811
    login(client, "U-B")
    assert "A IA está desligada" in client.get("/sugestoes").text


def test_linha_nova_da_planilha_sem_none_nas_novidades(client):  # noqa: F811
    folder = os.environ["LOCAL_FOLDER"]
    from pathlib import Path
    _editar_planilha(Path(folder), lambda ws, h: ws.append(
        ["ACT-109", "Preparar oficina de prompts", "Carla", datetime(2026, 10, 20), "Formação"]))
    login(client, "U-C")
    client.post("/sincronizar")
    texto = re.sub(r"<[^>]+>", " ", client.get("/novidades?desde=30d").text)
    assert "Preparar oficina de prompts" in texto and "None" not in texto


def test_precisa_de_conferencia_aparece_na_tela(client):  # noqa: F811
    c = _db()
    fid = c.execute("SELECT file_id FROM sources WHERE name='Ata_2026-10-04.md'").fetchone()[0]
    ver = c.execute("SELECT content_hash FROM sources WHERE file_id=?", (fid,)).fetchone()[0]
    c.execute("INSERT INTO extraction_notes (file_id, source_version, kind, text, reason, created_at) VALUES "
              "(?,?,?,?,?,?)", (fid, ver, "precisa_conferir", "Davi entrega o checklist até sexta que vem.",
                                "ACT-102: O prazo 16/10/2026 não aparece escrito como data no documento.", "2026-10-07T10:00:00"))
    c.commit()
    login(client, "U-B")
    pagina = client.get("/sugestoes").text
    assert "Precisa de conferência" in pagina and "sexta que vem" in pagina and "lo-conferir" in pagina


# ---------------------------------------------------------------------------
# Rodada 2 da avaliação da IA: situação dita no trecho, cancelamento e apelidos
# ---------------------------------------------------------------------------
def test_situacao_so_muda_quando_o_trecho_diz(conn, sync):
    sync()
    texto = ("Davi disse que ainda não conseguiu começar o checklist. Ficou combinado que ele entrega "
             "a primeira versão até 2026-10-16.")
    clean, problema = validate_item(conn, _item(target_activity_id="ACT-102", status="Em andamento",
                                                due_date="2026-10-16", evidence=texto), texto)
    assert problema is None and clean["proposed"] == {"due_date": "2026-10-16"}      # "ainda não começou" ≠ andamento
    assert any("não está dita no trecho" in u for u in clean["uncertainties"])


def test_bloqueada_volta_a_andar_quando_o_bloqueio_se_resolve(conn, sync):
    sync()
    texto = "Carla confirmou a sala da oficina."
    clean, problema = validate_item(conn, _item(target_activity_id="ACT-103", status="Em andamento",
                                                evidence=texto), texto)
    assert problema is None and clean["proposed"] == {"status": "Em andamento"}


def test_cancelamento_nunca_vira_concluida(conn, sync):
    sync()
    texto = "Decidimos não fazer mais a ACT-104; Ana e Davi ficam liberados."
    clean, problema = validate_item(conn, _item(target_activity_id="ACT-104", status="Concluída",
                                                evidence=texto), texto)
    assert problema is None and clean["proposed"] == {"status": "Cancelada"}
    assert any("cancelamento" in u for u in clean["uncertainties"])


def test_apelido_pede_conferencia(conn, sync):
    sync()
    texto = ("Participaram Ana, Bruno e Davi.\n\nO Bru falou que entra junto com a Aninha no carrossel (ACT-101). "
             "Novo prazo: 2026-10-09.")
    evid = "O Bru falou que entra junto com a Aninha no carrossel (ACT-101). Novo prazo: 2026-10-09."
    clean, problema = validate_item(conn, _item(target_activity_id="ACT-101", owners=["U-A", "U-B"],
                                                due_date="2026-10-09", evidence=evid), texto)
    assert problema is None and clean["proposed"]["owners"] == ["U-A", "U-B"]
    assert any("Bruno não aparece com esse nome no trecho" in u for u in clean["uncertainties"])


def test_prazo_fora_do_trecho_pede_conferencia(conn, sync):
    sync()
    texto = "Davi assume o checklist (ACT-102).\n\nOutro assunto: a reunião geral é em 2026-10-20."
    clean, _ = validate_item(conn, _item(target_activity_id="ACT-102", due_date="2026-10-20",
                                         evidence="Davi assume o checklist (ACT-102)."), texto)
    assert clean["proposed"]["due_date"] == "2026-10-20"
    assert any("não no trecho citado" in u for u in clean["uncertainties"])


def test_cancelar_pela_interface(client):  # noqa: F811
    login(client, "U-A")
    r = client.post("/atividades/ACT-104/estado", data={"status": "Cancelada", "reason": " "}, follow_redirects=False)
    assert "cancelar=1" in r.headers["location"]                                   # motivo obrigatório
    client.post("/atividades/ACT-104/estado", data={"status": "Cancelada", "reason": "O departamento troca de sistema"})
    c = _db()
    assert acts.snapshot(c, "ACT-104")["status"] == "Cancelada"
    assert "ACT-104" not in [a["activity_id"] for a in acts.list_activities(c, "U-A")]   # sai das abertas
    pagina = client.get("/atividades/ACT-104").text
    assert "Motivo do cancelamento" in pagina and "O departamento troca de sistema" in pagina and "Reabrir" in pagina
    assert "ACT-104" in client.get("/atividades?situacao=Cancelada").text
    valor = re.search(r'name="status" value="([^"]+)"><button[^>]*>Reabrir', pagina).group(1)
    client.post("/atividades/ACT-104/estado", data={"status": valor})
    assert acts.snapshot(_db(), "ACT-104")["status"] == "A fazer"
