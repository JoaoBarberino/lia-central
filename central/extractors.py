"""Extratores por formato: transformam bytes em texto/tabelas + metadados.

Cada formato tem seu extrator. Formatos fora do escopo retornam
`Unsupported` com o motivo, para a interface mostrar "ainda não processado".
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime

from .sources import CSV, DOCX, GDOC, GSHEET, GSLIDES, MARKDOWN, PDF, PPTX, TEXT, XLSX

REGISTER_HEADERS = ["ID", "Atividade", "Responsáveis", "Prazo", "Frente", "Prioridade", "Status",
                    "Próximo passo", "Origem", "Notas e bloqueios"]

META_LINE = re.compile(r"^\s*([a-z_]+)\s*:\s*(.+?)\s*$")


class Unsupported(Exception):
    """Formato fora do escopo implementado (não é erro de leitura)."""


class ExtractionError(Exception):
    """Arquivo no formato esperado, mas ilegível/corrompido."""


@dataclass
class Extracted:
    kind: str                      # "text" ou "table"
    title: str | None = None
    text: str = ""
    meta: dict = field(default_factory=dict)
    sheets: list[dict] = field(default_factory=list)  # [{name, headers, rows:[{row, cells:{col:valor}}]}]

    def canonical(self) -> str:
        """Representação estável usada para o hash de conteúdo (idempotência)."""
        if self.kind == "text":
            return self.text
        import json
        return json.dumps(self.sheets, ensure_ascii=False, sort_keys=True, default=str)


def normalize(s: str) -> str:
    """Minúsculas e sem acento: 'Responsáveis' -> 'responsaveis'."""
    s = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in s if not unicodedata.combining(c)).lower().strip()


# ---------------------------------------------------------------------------
# Texto (Markdown e Google Docs exportado como text/plain)
# ---------------------------------------------------------------------------
def parse_text_document(text: str) -> Extracted:
    # Google Docs exportado pode trazer quebras "moles" (\x0b, \u2028) e espaço não separável
    text = (text.replace("\ufeff", "").replace("\r\n", "\n").replace("\r", "\n")
            .replace("\x0b", "\n").replace("\u2028", "\n").replace("\u2029", "\n").replace("\xa0", " "))
    lines = text.split("\n")
    title = None
    meta: dict[str, str] = {}
    # Cabeçalho: título na primeira linha não vazia; depois linhas "chave: valor"
    # (linhas em branco entre elas são toleradas) até o primeiro parágrafo comum.
    for line in lines[:20]:
        stripped = line.strip()
        if not stripped:
            continue
        if title is None:
            title = stripped.lstrip("#").strip()
            continue
        m = META_LINE.match(stripped)
        if m:
            meta[m.group(1)] = m.group(2).strip().strip("`")
        else:
            break
    return Extracted(kind="text", title=title, text=text, meta=meta)


# ---------------------------------------------------------------------------
# Planilha .xlsx
# ---------------------------------------------------------------------------
def _cell_value(v):
    if isinstance(v, datetime):
        return v.date().isoformat()  # prazo: guardamos só a data
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return None
    if isinstance(v, str):
        return v.strip() or None
    return v


def parse_xlsx(data: bytes) -> Extracted:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # arquivo corrompido ou não é xlsx
        raise ExtractionError(f"Não foi possível abrir a planilha: {e}") from e
    sheets = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            sheets.append({"name": ws.title, "headers": [], "rows": []})
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        data_rows = []
        for i, r in enumerate(rows[1:], start=2):
            cells = {headers[j]: _cell_value(v) for j, v in enumerate(r) if j < len(headers) and headers[j]}
            if any(v is not None for v in cells.values()):
                data_rows.append({"row": i, "cells": cells})
        sheets.append({"name": ws.title, "headers": headers, "rows": data_rows})
    wb.close()
    return Extracted(kind="table", sheets=sheets)


def looks_like_register(sheet: dict) -> bool:
    """Uma aba 'parece registro de atividades' se tem as colunas-chave."""
    heads = {normalize(h) for h in sheet.get("headers", [])}
    return {"id", "atividade", "responsaveis"} <= heads


# ---------------------------------------------------------------------------
# PDF com texto selecionável (diferencial)
# ---------------------------------------------------------------------------
def parse_pdf(data: bytes) -> Extracted:
    import pdfplumber

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(p.extract_text() or "") for p in pdf.pages]
    except Exception as e:
        raise ExtractionError(f"Não foi possível abrir o PDF: {e}") from e
    text = "\n\n".join(pages).strip()
    if not text:
        raise Unsupported("PDF escaneado, sem texto para ler. O protótipo não lê texto dentro de imagens.")
    return parse_text_document(text)


# ---------------------------------------------------------------------------
# Word (.docx) e PowerPoint (.pptx): são pacotes zip com XML dentro. Lemos o texto
# diretamente (sem IA e sem biblioteca extra), na ordem do documento.
# ---------------------------------------------------------------------------
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _open_zip_xml(data: bytes, member: str, label: str):
    import zipfile
    import xml.etree.ElementTree as ET
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return ET.fromstring(z.read(member))
    except KeyError as e:
        raise ExtractionError(f"O arquivo não tem o conteúdo esperado de um {label}.") from e
    except Exception as e:  # zip corrompido, protegido por senha, XML inválido
        raise ExtractionError(f"Não foi possível abrir o {label}: {type(e).__name__}.") from e


def _docx_paragraph(p) -> str:
    out = []
    for el in p.iter():
        if el.tag == _W + "t":
            out.append(el.text or "")
        elif el.tag == _W + "tab":
            out.append("\t")
        elif el.tag in (_W + "br", _W + "cr"):
            out.append("\n")
    return "".join(out)


def parse_docx(data: bytes) -> Extracted:
    root = _open_zip_xml(data, "word/document.xml", "documento do Word")
    body = root.find(_W + "body")
    lines: list[str] = []
    for block in (list(body) if body is not None else []):
        if block.tag == _W + "p":
            lines.append(_docx_paragraph(block))
        elif block.tag == _W + "tbl":   # tabela: uma linha de texto por linha da tabela
            for tr in block.iter(_W + "tr"):
                cells = [" ".join(_docx_paragraph(p) for p in tc.iter(_W + "p")).strip() for tc in tr.iter(_W + "tc")]
                lines.append(" | ".join(cells))
    text = "\n".join(lines).strip()
    if not text:
        raise ExtractionError("O documento do Word está vazio (nenhum texto encontrado).")
    return parse_text_document(text)


def parse_pptx(data: bytes) -> Extracted:
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                            key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))
    except Exception as e:
        raise ExtractionError(f"Não foi possível abrir a apresentação: {type(e).__name__}.") from e
    parts = []
    for i, member in enumerate(slides, 1):
        root = _open_zip_xml(data, member, "apresentação")
        paras = ["".join(t.text or "" for t in p.iter(_A + "t")) for p in root.iter(_A + "p")]
        body = "\n".join(x for x in paras if x.strip())
        if body:
            parts.append(f"Slide {i}\n{body}")
    text = "\n\n".join(parts).strip()
    if not text:
        raise ExtractionError("A apresentação não tem texto (só imagens ou está vazia).")
    return parse_text_document(text)


# ---------------------------------------------------------------------------
# CSV: vira uma planilha de uma aba (mesmas regras do .xlsx, inclusive "planilha parecida com a oficial")
# ---------------------------------------------------------------------------
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def parse_csv(data: bytes, name: str = "dados") -> Extracted:
    import csv
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if text[:4096].count(";") > text[:4096].count(",") else ","
    rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ExtractionError("O arquivo CSV está vazio.")
    headers = [h.strip() for h in rows[0]]
    data_rows = []
    for i, r in enumerate(rows[1:], start=2):
        cells = {headers[j]: (v.strip() or None) for j, v in enumerate(r) if j < len(headers) and headers[j]}
        if any(v is not None for v in cells.values()):
            data_rows.append({"row": i, "cells": cells})
    sheet = name.rsplit(".", 1)[0] or "dados"
    return Extracted(kind="table", sheets=[{"name": sheet, "headers": headers, "rows": data_rows}])


# ---------------------------------------------------------------------------
IMAGE_MIMES = ("image/png", "image/jpeg", "image/webp", "image/heic", "image/heif", "image/gif")


def extract(mime_type: str, name: str, data: bytes) -> Extracted:
    low = name.lower()
    if mime_type in (MARKDOWN, GDOC, GSLIDES, TEXT) or low.endswith((".md", ".markdown", ".txt")):
        return parse_text_document(_decode(data))
    if mime_type in (XLSX, GSHEET) or low.endswith(".xlsx"):
        return parse_xlsx(data)
    if mime_type == CSV or low.endswith(".csv"):
        return parse_csv(data, name)
    if mime_type == PDF or low.endswith(".pdf"):
        return parse_pdf(data)
    if mime_type == DOCX or low.endswith(".docx"):
        return parse_docx(data)
    if mime_type == PPTX or low.endswith(".pptx"):
        return parse_pptx(data)
    if mime_type in IMAGE_MIMES or low.endswith((".png", ".jpg", ".jpeg", ".webp", ".heic")):
        raise Unsupported("Imagem: o protótipo não lê texto dentro de imagens automaticamente.")
    if low.endswith(".doc"):
        raise Unsupported("Arquivo .doc (Word antigo) não é lido. Salve como .docx ou como Documentos Google.")
    if low.endswith(".ppt"):
        raise Unsupported("Arquivo .ppt (PowerPoint antigo) não é lido. Salve como .pptx ou como Apresentações Google.")
    raise Unsupported(f"O protótipo ainda não lê este formato ({mime_type}).")


SUPPORTED_HINT = (".md, .txt, .docx, .pptx, .xlsx, .csv, PDF com texto, Google Docs, Google Sheets "
                  "e Google Slides")
